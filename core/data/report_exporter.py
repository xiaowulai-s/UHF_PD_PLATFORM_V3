# -*- coding: utf-8 -*-
"""PRPD 报告导出 — JSON/Excel/Word"""
from __future__ import annotations

import json
import os
from typing import Optional

from core.data.discharge_classifier import DischargeClassifier


class ReportExporter:
    """报告导出"""

    @staticmethod
    def export_json(data: dict, path: str) -> str:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2, default=str)
        return path

    @staticmethod
    def export_excel(data: dict, path: str) -> str:
        try:
            import openpyxl
        except ImportError:
            raise ImportError("需要安装 openpyxl: pip install openpyxl")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PRPD分析报告"

        # 标题
        ws.merge_cells("A1:E1")
        ws["A1"] = "PRPD 分析报告"
        ws["A1"].font = openpyxl.styles.Font(size=14, bold=True)

        # 基本信息
        ws["A3"] = "分析时间"
        ws["B3"] = data.get("analysis_time", "")
        ws["A4"] = "放电类型(参考结果)"
        dc = data.get("discharge_classification", {})
        ref = dc.get("reference_result", {}) if isinstance(dc, dict) else {}
        ws["B4"] = ref.get("name", dc.get("name", "") if isinstance(dc, dict) else "")
        ws["A5"] = "置信度(占比)"
        ws["B5"] = f'{ref.get("confidence", dc.get("confidence", 0)):.0%}' if isinstance(dc, dict) else "0%"
        ws["A6"] = "严重等级"
        ws["B6"] = dc.get("severity", "") if isinstance(dc, dict) else ""
        ws["A7"] = "总周期数"
        ws["B7"] = str(dc.get("total_samples", "")) if isinstance(dc, dict) else ""

        # 统计
        stats = data.get("statistics", {})
        ws["A9"] = "统计指标"
        ws["A9"].font = openpyxl.styles.Font(bold=True)
        ws["A10"] = "总事件数"
        ws["B10"] = stats.get("total_events", 0)
        ws["A11"] = "最大幅值"
        ws["B11"] = stats.get("max_amplitude", 0)
        ws["A12"] = "平均幅值"
        ws["B12"] = stats.get("avg_amplitude", 0)
        ws["A13"] = "幅值中位数"
        ws["B13"] = stats.get("median_amplitude", 0)
        ws["A14"] = "幅值P90"
        ws["B14"] = stats.get("p90", 0)
        ws["A15"] = "最小幅值"
        ws["B15"] = stats.get("min_amplitude", 0)

        # 各类型统计
        if isinstance(dc, dict):
            type_counts = dc.get("type_counts", {})
            type_ratios = dc.get("type_ratios", {})
            row = 17
            ws[f"A{row}"] = "放电类型统计"
            ws[f"A{row}"].font = openpyxl.styles.Font(bold=True)
            row += 1
            for t, cnt in sorted(type_counts.items(), key=lambda x: x[1], reverse=True):
                pattern = DischargeClassifier.PATTERNS.get(t, {})
                ws[f"A{row}"] = pattern.get("name", t)
                ws[f"B{row}"] = str(cnt)
                ws[f"C{row}"] = f"{type_ratios.get(t, 0):.1%}"
                row += 1

        wb.save(path)
        return path

    @staticmethod
    def export_word(data: dict, path: str, chart_path: Optional[str] = None) -> str:
        """导出 Word 报告（含 PRPD 截图）"""
        # 图表存在性先算，再导 docx。放在导库之后会有个副作用：这台机器没装
        # python-docx，而用例普遍套着 importorskip("docx")，于是带 chart_path 的
        # 分支从来没被执行过——正是这个模块缺 `import os` 却长期全绿的原因。
        include_chart = bool(chart_path) and os.path.exists(chart_path)

        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches
        except ImportError:
            raise ImportError("需要安装 python-docx: pip install python-docx")

        doc = Document()

        # 标题
        title = doc.add_heading("PRPD 局部放电分析报告", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 基本信息
        doc.add_heading("1. 基本信息", level=1)
        tbl = doc.add_table(rows=5, cols=2)
        tbl.style = "Light Shading Accent 1"
        items = [
            ("分析时间", data.get("analysis_time", "")),
            ("数据文件", data.get("file_name", "")),
            ("总事件数", str(data.get("statistics", {}).get("total_events", 0))),
            ("放电类型", data.get("discharge_classification", {}).get("name", "")),
            ("严重等级", data.get("discharge_classification", {}).get("severity", "")),
        ]
        for i, (k, v) in enumerate(items):
            tbl.cell(i, 0).text = k
            tbl.cell(i, 1).text = v

        # PRPD 图谱截图
        if include_chart:
            doc.add_heading("2. PRPD 图谱", level=1)
            doc.add_picture(chart_path, width=Inches(5.5))

        # 分类结果
        dc = data.get("discharge_classification", {})
        doc.add_heading("3. 放电类型分析（按周期统计）", level=1)
        if isinstance(dc, dict):
            ref = dc.get("reference_result", {})
            doc.add_paragraph(
                f"参考结果: {ref.get('name', dc.get('name', ''))} ({ref.get('name_en', '')}) [标记为参考结果]"
            )
            doc.add_paragraph(f"置信度(占比): {ref.get('confidence', dc.get('confidence', 0)):.0%}")
            doc.add_paragraph(f"采样次数: {dc.get('total_samples', 0)}")
            doc.add_paragraph(f"严重等级: {dc.get('severity', 'normal')}")

            # 各类型统计表格
            type_counts = dc.get("type_counts", {})
            type_ratios = dc.get("type_ratios", {})
            if type_counts:
                tbl3 = doc.add_table(rows=len(type_counts) + 1, cols=3)
                tbl3.style = "Light Shading Accent 1"
                tbl3.cell(0, 0).text = "放电类型"
                tbl3.cell(0, 1).text = "出现次数"
                tbl3.cell(0, 2).text = "占比"
                for i, (t, cnt) in enumerate(sorted(type_counts.items(), key=lambda x: x[1], reverse=True), 1):
                    pattern = DischargeClassifier.PATTERNS.get(t, {})
                    tbl3.cell(i, 0).text = pattern.get("name", t)
                    tbl3.cell(i, 1).text = str(cnt)
                    tbl3.cell(i, 2).text = f"{type_ratios.get(t, 0):.1%}"
        else:
            doc.add_paragraph(f"识别结果: {dc.get('name', '')}")
            doc.add_paragraph(f"置信度: {dc.get('confidence', 0):.0%}")

        # 统计指标
        stats = data.get("statistics", {})
        doc.add_heading("4. 统计指标", level=1)
        tbl2 = doc.add_table(rows=8, cols=2)
        tbl2.style = "Light Shading Accent 1"
        stat_items = [
            ("总事件数", str(stats.get("total_events", 0))),
            ("最大幅值", f'{stats.get("max_amplitude", 0):.1f}'),
            ("最小幅值", f'{stats.get("min_amplitude", 0):.1f}'),
            ("平均幅值", f'{stats.get("avg_amplitude", 0):.1f}'),
            ("幅值中位数", f'{stats.get("median_amplitude", 0):.1f}'),
            ("幅值P90", f'{stats.get("p90", 0):.1f}'),
        ]
        for i, (k, v) in enumerate(stat_items):
            tbl2.cell(i, 0).text = k
            tbl2.cell(i, 1).text = v

        doc.save(path)
        return path


# ═══════════════════════════════════════════════════════
# 数据分析页面
# ═══════════════════════════════════════════════════════

# 等级颜色 (硬编码避免 core 层依赖 UI 层)
LEVEL_COLORS = {
    "critical": "#DC2626",
    "high": "#EA580C",
    "warning": "#D29922",
    "medium": "#CA8A04",
    "attention": "#2563EB",
    "low": "#2563EB",
    "normal": "#16A34A",
}
LEVEL_NAMES = {
    "critical": "严重",
    "high": "高级",
    "warning": "警告",
    "medium": "中级",
    "low": "低级",
    "attention": "注意",
    "normal": "正常",
}
