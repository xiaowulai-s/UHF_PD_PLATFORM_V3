# -*- coding: utf-8 -*-
"""PRPD 实验数据导入器 — CSV/Excel/TXT 自动识别"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Optional, Tuple

import numpy as np


class DataImporter:
    """PRPD 实验数据导入器"""

    SUPPORTED_EXTENSIONS = [".csv", ".xlsx", ".xls", ".txt", ".npy", ".npz", ".h5", ".hdf5", ".json"]

    @staticmethod
    def load(path: str) -> Tuple[str, np.ndarray, np.ndarray, np.ndarray, dict]:
        """
        加载数据文件

        Returns:
            (format_name, phases, amplitudes, cycles, metadata)
        """
        suffix = Path(path).suffix.lower()
        if suffix not in DataImporter.SUPPORTED_EXTENSIONS:
            raise ValueError(f"不支持的文件格式: {suffix}，支持: {', '.join(DataImporter.SUPPORTED_EXTENSIONS)}")

        if suffix == ".csv":
            return DataImporter._load_csv(path)
        elif suffix in (".xlsx", ".xls"):
            return DataImporter._load_excel(path)
        elif suffix in (".npy", ".npz"):
            return DataImporter._load_numpy(path)
        elif suffix in (".h5", ".hdf5"):
            return DataImporter._load_hdf5(path)
        elif suffix == ".json":
            return DataImporter._load_json(path)
        else:
            return DataImporter._load_txt(path)

    @staticmethod
    def _load_csv(path: str) -> Tuple[str, np.ndarray, np.ndarray, np.ndarray, dict]:
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = [r for r in reader if r and not r[0].startswith(("#", "%", "//"))]

        if not rows:
            raise ValueError("文件为空或不包含有效数据行")

        header = [h.strip().lower() for h in rows[0]]
        data_rows = rows[1:]

        # 自动映射列
        phase_idx, amp_idx, cycle_idx = None, None, None
        for i, h in enumerate(header):
            if h in ("phase_deg", "phase", "phi", "φ", "相位"):
                phase_idx = i
            if h in ("discharge", "amplitude", "amp", "magnitude", "幅值", "mv", "pc", "放电量"):
                amp_idx = i
            if h in ("cycle", "周期", "n", "period", "工频周期"):
                cycle_idx = i

        phases, amplitudes, cycles = [], [], []
        for row in data_rows:
            try:
                p = float(row[phase_idx]) if phase_idx is not None and len(row) > phase_idx else None
                a = float(row[amp_idx]) if amp_idx is not None and len(row) > amp_idx else None
                c = int(row[cycle_idx]) if cycle_idx is not None and len(row) > cycle_idx else 0
                if p is not None and a is not None and 0 <= p <= 360 and a >= 0:
                    phases.append(p)
                    amplitudes.append(a)
                    cycles.append(c)
            except (ValueError, IndexError, TypeError):
                continue

        if not phases:
            # 回退: 尝试前两列
            for row in data_rows:
                if len(row) >= 2:
                    try:
                        p, a = float(row[0]), float(row[1])
                        if 0 <= p <= 360 and a >= 0:
                            phases.append(p)
                            amplitudes.append(a)
                            cycles.append(0)
                    except (ValueError, IndexError):
                        continue

        return (
            "phase_amplitude",
            np.array(phases),
            np.array(amplitudes),
            np.array(cycles, dtype=np.int32),
            {
                "columns": header,
                "total_rows": len(data_rows),
                "file": path,
            },
        )

    @staticmethod
    def _load_excel(path: str) -> Tuple[str, np.ndarray, np.ndarray, np.ndarray, dict]:
        try:
            import openpyxl
        except ImportError:
            raise ImportError("需要安装 openpyxl: pip install openpyxl")

        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        header = [str(c).strip().lower() if c else "" for c in all_rows[0]]
        data_rows = all_rows[1:]
        wb.close()

        phase_idx, amp_idx, cycle_idx = None, None, None
        for i, h in enumerate(header):
            if h in ("phase_deg", "phase", "phi", "φ", "相位"):
                phase_idx = i
            if h in ("discharge", "amplitude", "amp", "magnitude", "幅值", "mv", "pc"):
                amp_idx = i
            if h in ("cycle", "周期", "n", "period"):
                cycle_idx = i

        phases, amplitudes, cycles = [], [], []
        default_idx = 0
        for row in data_rows:
            try:
                p = (
                    float(row[phase_idx])
                    if phase_idx is not None and len(row) > phase_idx
                    else float(row[default_idx]) if row[default_idx] is not None else None
                )
                a = (
                    float(row[amp_idx])
                    if amp_idx is not None and len(row) > amp_idx
                    else (
                        float(row[default_idx + 1])
                        if len(row) > default_idx + 1 and row[default_idx + 1] is not None
                        else None
                    )
                )
                c = int(row[cycle_idx]) if cycle_idx is not None and len(row) > cycle_idx else 0
                if p is not None and a is not None and 0 <= p <= 360 and a >= 0:
                    phases.append(p)
                    amplitudes.append(a)
                    cycles.append(c)
            except (ValueError, IndexError, TypeError):
                continue

        return (
            "phase_amplitude",
            np.array(phases),
            np.array(amplitudes),
            np.array(cycles, dtype=np.int32),
            {
                "columns": header,
                "total_rows": len(data_rows),
                "file": path,
            },
        )

    @staticmethod
    def _load_txt(path: str) -> Tuple[str, np.ndarray, np.ndarray, np.ndarray, dict]:
        with open(path, "r", encoding="utf-8") as f:
            lines = [ln.strip() for ln in f if ln.strip() and not ln.strip().startswith(("#", "%"))]

        phases, amplitudes, cycles = [], [], []
        for line in lines:
            parts = line.split()
            if len(parts) >= 2:
                try:
                    p, a = float(parts[0]), float(parts[1])
                    if 0 <= p <= 360:
                        phases.append(p)
                        amplitudes.append(a)
                        cycles.append(0)
                except ValueError:
                    continue

        return (
            "phase_amplitude",
            np.array(phases),
            np.array(amplitudes),
            np.array(cycles, dtype=np.int32),
            {
                "columns": ["col1", "col2"],
                "total_rows": len(lines),
                "file": path,
            },
        )

    @staticmethod
    def _load_numpy(path: str) -> Tuple[str, np.ndarray, np.ndarray, np.ndarray, dict]:
        """加载 NumPy 二进制格式 (.npy/.npz)"""
        try:
            data = np.load(path, allow_pickle=True)
        except ImportError:
            raise ImportError("需要安装 numpy: pip install numpy")

        # 支持多种数据结构
        if isinstance(data, np.lib.npyio.NpzFile):
            # .npz 文件：尝试自动检测键
            if "phases" in data and "amplitudes" in data:
                phases = data["phases"]
                amplitudes = data["amplitudes"]
                cycles = data.get("cycles", np.zeros(len(phases), dtype=np.int32))
            elif "phase" in data and "amplitude" in data:
                phases = data["phase"]
                amplitudes = data["amplitude"]
                cycles = data.get("cycle", np.zeros(len(phases), dtype=np.int32))
            else:
                # 回退：使用前两个数组
                keys = list(data.keys())
                if len(keys) >= 2:
                    phases = data[keys[0]].flatten()
                    amplitudes = data[keys[1]].flatten()
                    cycles = np.zeros(len(phases), dtype=np.int32)
                else:
                    raise ValueError("无法识别的 .npz 数据结构")
        else:
            # .npy 文件：假设是 shape (N, 2) 或 (N, 3) 的数组
            arr = np.atleast_2d(data)
            if arr.shape[1] >= 2:
                phases = arr[:, 0]
                amplitudes = arr[:, 1]
                cycles = arr[:, 2] if arr.shape[1] >= 3 else np.zeros(len(phases), dtype=np.int32)
            else:
                raise ValueError(f".npy 文件需要至少 2 列数据，实际形状: {arr.shape}")

        metadata = {
            "columns": ["phase", "amplitude", "cycle"],
            "total_rows": len(phases),
            "file": path,
            "format": "numpy",
        }

        return ("phase_amplitude", phases, amplitudes, cycles, metadata)

    @staticmethod
    def _load_hdf5(path: str) -> Tuple[str, np.ndarray, np.ndarray, np.ndarray, dict]:
        """加载 HDF5 格式 (.h5/.hdf5)"""
        try:
            import h5py
        except ImportError:
            raise ImportError("需要安装 h5py: pip install h5py")

        with h5py.File(path, "r") as f:
            # 自动检测数据集
            if "phases" in f and "amplitudes" in f:
                phases = f["phases"][:]
                amplitudes = f["amplitudes"][:]
                cycles = f.get("cycles", np.zeros(len(phases), dtype=np.int32))[:]
            elif "phase" in f and "amplitude" in f:
                phases = f["phase"][:]
                amplitudes = f["amplitude"][:]
                cycles = f.get("cycle", np.zeros(len(phases), dtype=np.int32))[:]
            else:
                # 回退：使用前两个数据集
                keys = list(f.keys())
                if len(keys) >= 2:
                    phases = f[keys[0]][:].flatten()
                    amplitudes = f[keys[1]][:].flatten()
                    cycles = np.zeros(len(phases), dtype=np.int32)
                else:
                    raise ValueError("无法识别的 HDF5 数据结构")

            metadata = {
                "columns": list(f.keys()),
                "total_rows": len(phases),
                "file": path,
                "format": "hdf5",
            }

        return ("phase_amplitude", phases, amplitudes, cycles, metadata)

    @staticmethod
    def _load_json(path: str) -> Tuple[str, np.ndarray, np.ndarray, np.ndarray, dict]:
        """加载 JSON 格式"""
        with open(path, "r", encoding="utf-8") as f:
            import json

            data = json.load(f)

        # 支持多种 JSON 结构
        if isinstance(data, list):
            # [{phase: ..., amplitude: ...}, ...]
            phases, amplitudes, cycles = [], [], []
            for item in data:
                if isinstance(item, dict):
                    p = item.get("phase", item.get("phase_deg"))
                    a = item.get("amplitude", item.get("discharge", item.get("amp")))
                    c = item.get("cycle", 0)
                    if p is not None and a is not None:
                        phases.append(float(p))
                        amplitudes.append(float(a))
                        cycles.append(int(c))
        elif isinstance(data, dict):
            # {"phases": [...], "amplitudes": [...], ...}
            phases = data.get("phases", data.get("phase", []))
            amplitudes = data.get("amplitudes", data.get("amplitude", []))
            cycles = data.get("cycles", data.get("cycle", [0] * len(phases)))

        if not phases:
            raise ValueError("JSON 文件中未找到有效数据")

        metadata = {
            "columns": ["phase", "amplitude", "cycle"],
            "total_rows": len(phases),
            "file": path,
            "format": "json",
        }

        return ("phase_amplitude", np.array(phases), np.array(amplitudes), np.array(cycles, dtype=np.int32), metadata)

    @staticmethod
    def validate_data(phases: np.ndarray, amplitudes: np.ndarray, cycles: Optional[np.ndarray] = None) -> dict:
        """
        验证数据质量并生成质量报告

        Args:
            phases: 相位数组 (0~360°)
            amplitudes: 幅值数组
            cycles: 周期数组（可选）

        Returns:
            质量报告字典
        """
        report = {
            "valid": True,
            "total_events": len(phases),
            "valid_events": 0,
            "invalid_events": 0,
            "phase_range": (
                float(np.min(phases)) if len(phases) > 0 else 0,
                float(np.max(phases)) if len(phases) > 0 else 0,
            ),
            "amplitude_range": (
                float(np.min(amplitudes)) if len(amplitudes) > 0 else 0,
                float(np.max(amplitudes)) if len(amplitudes) > 0 else 0,
            ),
            "duplicates": 0,
            "outliers": 0,
            "warnings": [],
        }

        if len(phases) == 0:
            report["valid"] = False
            report["warnings"].append("数据为空")
            return report

        invalid_phase = (phases < 0) | (phases > 360)
        invalid_amp = amplitudes < 0
        invalid_mask = invalid_phase | invalid_amp
        report["invalid_events"] = int(np.sum(invalid_mask))
        report["valid_events"] = len(phases) - report["invalid_events"]

        if len(phases) > 1:
            combined = np.column_stack((phases, amplitudes))
            unique_rows = np.unique(combined, axis=0)
            report["duplicates"] = len(combined) - len(unique_rows)

        if len(amplitudes) > 4:
            q1 = np.percentile(amplitudes, 25)
            q3 = np.percentile(amplitudes, 75)
            iqr = q3 - q1
            outliers = amplitudes[(amplitudes < q1 - 1.5 * iqr) | (amplitudes > q3 + 1.5 * iqr)]
            report["outliers"] = len(outliers)

        if report["invalid_events"] > len(phases) * 0.1:
            report["valid"] = False
            report["warnings"].append("无效事件过多")

        return report

    @staticmethod
    def preview_file(path: str, max_rows: int = 10) -> dict:
        """预览文件内容"""
        suffix = Path(path).suffix.lower()

        if suffix == ".csv":
            with open(path, "r", encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                rows = []
                for i, row in enumerate(reader):
                    if i >= max_rows + 1:
                        break
                    rows.append(row)

            header = rows[0] if rows else []
            preview = rows[1 : max_rows + 1]
            return {
                "format": "CSV",
                "total_rows": "未知",
                "columns": header,
                "preview": preview,
            }

        elif suffix in (".xlsx", ".xls"):
            try:
                import openpyxl

                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                ws = wb.active
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= max_rows + 1:
                        break
                    rows.append(list(row))
                wb.close()

                header = rows[0] if rows else []
                preview = rows[1 : max_rows + 1]
                return {
                    "format": "Excel",
                    "total_rows": "未知",
                    "columns": [str(c) if c else "" for c in header],
                    "preview": preview,
                }
            except ImportError:
                raise ImportError("需要安装 openpyxl")

        elif suffix in (".npy", ".npz"):
            data = np.load(path, allow_pickle=True)
            if isinstance(data, np.lib.npyio.NpzFile):
                shape = {k: data[k].shape for k in data.keys()}
                preview_data = {k: data[k][:5].tolist() for k in list(data.keys())[:3]}
                columns = list(data.keys())
            else:
                shape = data.shape
                preview_data = np.atleast_2d(data)[:5].tolist()
                columns = ["col1", "col2"]

            return {
                "format": "NumPy",
                "total_rows": str(shape),
                "columns": columns,
                "preview": preview_data,
            }

        elif suffix in (".h5", ".hdf5"):
            try:
                import h5py

                with h5py.File(path, "r") as f:
                    keys = list(f.keys())
                    shape = {k: f[k].shape for k in keys[:3]}
                    preview_data = {k: f[k][:5].tolist() for k in keys[:3]}

                return {
                    "format": "HDF5",
                    "total_rows": str(shape),
                    "columns": keys,
                    "preview": preview_data,
                }
            except ImportError:
                raise ImportError("需要安装 h5py")

        elif suffix == ".json":
            with open(path, "r", encoding="utf-8") as f:
                import json

                data = json.load(f)

            if isinstance(data, list):
                preview = data[:max_rows]
                total = len(data)
                columns = list(data[0].keys()) if data and isinstance(data[0], dict) else []
            elif isinstance(data, dict):
                preview = {k: v for i, (k, v) in enumerate(data.items()) if i < max_rows}
                total = len(data)
                columns = list(data.keys())
            else:
                preview = [str(data)]
                total = 1
                columns = []

            return {
                "format": "JSON",
                "total_rows": total,
                "columns": columns,
                "preview": preview,
            }

        else:
            raise ValueError(f"不支持的文件格式: {suffix}")


# ═══════════════════════════════════════════════════════
# ═══════════════════════════════════════════════════════
