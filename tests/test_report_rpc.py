# -*- coding: utf-8 -*-
"""report.export RPC 的输入契约与产出内容。

这里刻意不重复测 ReportExporter 本身（tests/test_report_exporter.py 已经按路径测过）。
本文件钉的是 RPC 这一层：格式白名单、data 的形状、图表载荷的解码失败必须报错而不是
静默出一张没图的报告，以及"返回的字节真的是一份能重新打开的报告"。
"""

import base64
import io
import json

import pytest

from core.protocol.report_rpc import ReportRpc, register_report_rpc

#: 1x1 透明 PNG，够让 python-docx 认成图片又不必依赖渲染端
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


@pytest.fixture
def sample():
    return {
        "analysis_time": "2026-10-10T09:00:00",
        "file_name": "PD_data.csv",
        "statistics": {
            "total_events": 6657,
            "max_amplitude": 8462.0,
            "min_amplitude": 12.5,
            "avg_amplitude": 320.5,
            "median_amplitude": 210.0,
            "p90": 700.0,
        },
        "discharge_classification": {
            "name": "内部放电",
            "confidence": 0.62,
            "severity": "high",
            "total_samples": 100,
            "type_counts": {"internal": 62, "corona": 38},
            "type_ratios": {"internal": 0.62, "corona": 0.38},
        },
    }


def test_json_report_is_parseable_and_keeps_the_numbers(sample):
    out = ReportRpc().export_report({"format": "json", "data": sample})

    assert out["format"] == "json"
    assert out["filename"].endswith(".json")
    loaded = json.loads(out["bytes"].decode("utf-8"))
    assert loaded["statistics"]["total_events"] == 6657
    assert loaded["discharge_classification"]["name"] == "内部放电"
    # 中文不能被转义成 \uXXXX 之外还能读——ensure_ascii=False 的可见证据
    assert "内部放电" in out["bytes"].decode("utf-8")
    assert out["size"] == len(out["bytes"])


def test_excel_report_reopens_and_shows_the_headline_values(sample):
    pytest.importorskip("openpyxl")
    import openpyxl

    out = ReportRpc().export_report({"format": "excel", "data": sample})

    wb = openpyxl.load_workbook(io.BytesIO(out["bytes"]))
    ws = wb.active
    cells = {c.value: wb.active.cell(c.row, c.column + 1).value for c in ws["A"] if c.value}
    assert ws["A1"].value == "PRPD 分析报告"
    assert cells["总事件数"] == 6657
    assert cells["最大幅值"] == 8462.0
    assert cells["放电类型(参考结果)"] == "内部放电"
    assert cells["严重等级"] == "high"


def test_word_report_without_chart_reopens(sample):
    pytest.importorskip("docx")
    import docx

    out = ReportRpc().export_report({"format": "word", "data": sample})

    doc = docx.Document(io.BytesIO(out["bytes"]))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "PRPD 局部放电分析报告" in text
    assert "3. 放电类型分析" in text
    assert out["chart_embedded"] is False


def test_word_report_with_chart_embeds_the_image(sample):
    """有图/无图两条路径都要真跑过：原先只有无图那条被覆盖，缺 import os 才藏得住。"""
    pytest.importorskip("docx")
    import docx

    out = ReportRpc().export_report(
        {"format": "word", "data": sample, "chart_png": base64.b64encode(PNG_1PX).decode("ascii")}
    )

    doc = docx.Document(io.BytesIO(out["bytes"]))
    assert out["chart_embedded"] is True
    assert "2. PRPD 图谱" in "\n".join(p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 1, "报告里必须真的嵌进了那张图"


def test_chart_accepts_raw_bytes_too(sample):
    """前端 dataURL 去掉前缀后可能是 Uint8Array，直接当 bin 传过来。"""
    out = ReportRpc().export_report({"format": "json", "data": sample, "chart_png": PNG_1PX})
    assert out["chart_embedded"] is False  # json 不需要图，但载荷不能被误拒


@pytest.mark.parametrize(
    "params, message",
    [
        ({"format": "pdf", "data": {"a": 1}}, "未知报告格式"),
        ({"format": "json", "data": "不是字典"}, "非空对象"),
        ({"format": "json", "data": {}}, "非空对象"),
        ({"data": {"a": 1}}, "未知报告格式"),
        ({"format": "json", "data": {"a": 1}, "chart_png": "这不是base64!!!"}, "不是合法 base64"),
        ({"format": "json", "data": {"a": 1}, "chart_png": "aGVsbG8="}, "不是 PNG"),
    ],
)
def test_bad_input_is_rejected_with_a_readable_message(params, message):
    with pytest.raises(ValueError, match=message):
        ReportRpc().export_report(params)


def test_oversized_data_is_rejected():
    big = {"blob": "x" * (9 * 1024 * 1024)}
    with pytest.raises(ValueError, match="超过上限"):
        ReportRpc().export_report({"format": "json", "data": big})


def test_extension_comes_from_the_whitelist_not_the_caller(sample):
    """前端传什么都不该影响落盘后缀——那才是会写错文件的地方。"""
    out = ReportRpc().export_report({"format": "excel", "data": sample, "filename": "../evil.py"})
    assert out["filename"].endswith(".xlsx")
    assert "/" not in out["filename"] and ".." not in out["filename"]


class _FakeServer:
    def __init__(self):
        self.registered = {}

    def register_rpc(self, method, handler, background=False):
        self.registered[method] = (handler, background)


def test_registration_is_background_so_generation_never_blocks_the_qt_loop():
    server = _FakeServer()

    assert register_report_rpc(server) == ["report.formats", "report.export"]
    for name in ("report.formats", "report.export"):
        handler, background = server.registered[name]
        assert background is True, "导出是 CPU+磁盘活，跑在 Qt 事件循环上会卡推流"
        assert callable(handler)


def test_formats_agrees_with_what_is_actually_importable():
    """按钮能不能点由内核自报，而这个自报必须和真实导入能力一致。

    不硬编码"本机没有 docx"——那会让装了 python-docx 的机器（或以后补依赖）红掉。
    """
    import importlib.util

    out = ReportRpc().available_formats({})

    assert "json" in out["available"]
    for fmt, mod in (("excel", "openpyxl"), ("word", "docx")):
        present = importlib.util.find_spec(mod) is not None
        assert (fmt in out["available"]) is present, f"{fmt} 的自报能力与 {mod} 是否存在不符"
        if not present:
            assert mod.split("_")[0].lower() in out["missing"][fmt].lower() or out["missing"][fmt]
