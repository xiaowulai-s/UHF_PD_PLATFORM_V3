"""P2 报告导出的真点验收（CDP 驱动，跑在打包后的 Electron 里）。

为什么不能只在 Python 里测：按钮要穿过 store.call → msgpack bin → preload 桥 →
主进程对话框这条完整链，任何一环没接上，Python 侧全绿也看不出来。

难点在"怎么让页面里出现一份导入数据"：Chromium 不允许脚本给 input[type=file] 赋值，
所以走 CDP 的 DOM.setFileInputFiles —— 它塞的是**真实磁盘路径**，preload 里的
webUtils.getPathForFile 才换得到内核能读的路径（这正是导入链的入口条件）。
点击用 Input.dispatchMouseEvent，并且按本仓库踩过的三条：先 mouseMoved、
clickCount 显式给 1、拖拽不算。

落盘用 PD_SAVE_PATH（只在 CAPTURE 模式下生效），拿到文件后用 openpyxl/json 重新打开，
而不是只看按钮文字变了没。用完即删。
"""

import argparse
import json
import os
import pathlib
import shutil
import socket
import subprocess
import sys
import time
import urllib.request

from websockets.sync.client import connect

REPO = pathlib.Path(__file__).resolve().parents[1]
SHOTS = REPO / "tmp" / "p2-verify"
PORT = 9351
CSV = REPO / "tests" / "PD_data.csv"
SAVE = SHOTS / "out.bin"

_id = 0


def rpc(ws, method, params=None):
    global _id
    _id += 1
    mine = _id
    ws.send(json.dumps({"id": mine, "method": method, "params": params or {}}))
    deadline = time.time() + 30
    while time.time() < deadline:
        body = json.loads(ws.recv(timeout=30))
        if body.get("id") != mine:
            continue
        if "error" in body:
            raise RuntimeError(f"{method}: {body['error'].get('message')}")
        return body.get("result", {})
    raise TimeoutError(method)


def evaluate(ws, expr):
    out = rpc(ws, "Runtime.evaluate", {"expression": expr, "returnByValue": True, "awaitPromise": True})
    exc = out.get("exceptionDetails")
    if exc:
        raise RuntimeError(f"页面表达式抛错：{json.dumps(exc)[:200]}")
    return (out.get("result") or {}).get("value")


def wait_for(ws, expr, what, timeout=60.0):
    end = time.time() + timeout
    while time.time() < end:
        try:
            val = evaluate(ws, expr)
        except Exception:
            val = None
        if val:
            return val
        time.sleep(0.4)
    raise TimeoutError(f"等不到：{what}")


def click_text(ws, label, tag="button"):
    rect = evaluate(
        ws,
        "(() => {const el=[...document.querySelectorAll(%s)].find(e=>e.textContent.trim()===%s);"
        "if(!el) return null; const r=el.getBoundingClientRect();"
        "return {x: Math.round(r.left+r.width/2), y: Math.round(r.top+r.height/2),"
        " disabled: !!el.disabled, title: el.title || ''};})()" % (json.dumps(tag), json.dumps(label)),
    )
    if not rect:
        raise AssertionError(f"找不到 {tag}: {label}")
    rpc(ws, "Input.dispatchMouseEvent", {"type": "mouseMoved", "x": rect["x"], "y": rect["y"]})
    rpc(
        ws,
        "Input.dispatchMouseEvent",
        {"type": "mousePressed", "x": rect["x"], "y": rect["y"], "button": "left", "clickCount": 1, "buttons": 1},
    )
    rpc(
        ws,
        "Input.dispatchMouseEvent",
        {"type": "mouseReleased", "x": rect["x"], "y": rect["y"], "button": "left", "clickCount": 1, "buttons": 0},
    )
    return rect


def parse_mode():
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--packaged",
        action="store_true",
        help="验 desktop/release/win-unpacked 里的装机版（内核来自 resources/uhf-sidecar），"
        "而不是开发态的 electron.exe + 源码 python",
    )
    return ap.parse_args([] if len(sys.argv) < 2 else sys.argv[1:])


def main():
    args = parse_mode()
    if not CSV.is_file():
        print(f"缺少样例数据 {CSV}")
        return 1
    SHOTS.mkdir(parents=True, exist_ok=True)
    for stale in SHOTS.glob("out*.bin"):
        stale.unlink()

    env = {k: v for k, v in os.environ.items() if k != "ELECTRON_RUN_AS_NODE"}
    if args.packaged:
        exe = str(REPO / "desktop" / "release" / "win-unpacked" / "UHF PD 监测系统.exe")
        launch = [exe, "--disable-gpu-sandbox", "--no-sandbox", f"--remote-debugging-port={PORT}"]
        cwd = str(REPO / "desktop" / "release" / "win-unpacked")
    else:
        exe = str(REPO / "desktop" / "node_modules" / "electron" / "dist" / "electron.exe")
        launch = [exe, ".", "--disable-gpu-sandbox", "--no-sandbox", f"--remote-debugging-port={PORT}"]
        cwd = str(REPO / "desktop")

    env.update(
        {
            "PD_ROUTE": "analysis",
            "PD_CAPTURE": str(SHOTS / "shot.png"),
            "PD_CAPTURE_WAIT": "240000",
            "PD_SAVE_PATH": str(SAVE),
            "PD_SIDECAR_ARGS": "--pd-type corona --pulse-rate 20000",
        }
    )
    proc = subprocess.Popen(
        launch,
        cwd=cwd,
        env=env,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    failures = []
    try:
        url = None
        end = time.time() + 60
        while time.time() < end and url is None:
            try:
                with socket.create_connection(("127.0.0.1", PORT), timeout=1):
                    with urllib.request.urlopen(f"http://127.0.0.1:{PORT}/json", timeout=2) as r:
                        for t in json.load(r):
                            if t.get("type") == "page" and t.get("webSocketDebuggerUrl"):
                                url = t["webSocketDebuggerUrl"]
                                break
            except OSError:
                time.sleep(0.5)
        if url is None:
            print("CDP 未就绪")
            return 1

        with connect(url, open_timeout=10, close_timeout=5, max_size=None) as ws:
            rpc(ws, "Runtime.enable")
            rpc(ws, "DOM.enable")
            wait_for(ws, "!!document.querySelector('input[type=file]')", "分析页的文件选择框")

            root = rpc(ws, "DOM.getDocument")["root"]["nodeId"]
            node = rpc(ws, "DOM.querySelector", {"nodeId": root, "selector": "input[type=file]"})["nodeId"]
            if not node:
                print("FAIL: 拿不到 input[type=file] 的 nodeId")
                return 1
            rpc(ws, "DOM.setFileInputFiles", {"files": [str(CSV)], "nodeId": node})
            evaluate(
                ws,
                "(()=>{const i=document.querySelector('input[type=file]');"
                "i.dispatchEvent(new Event('change',{bubbles:true}));})()",
            )
            print("已把真实路径交给 input[type=file] 并派发 change，等导入完成…")

            wait_for(ws, "/导出分析报告/.test(document.body.innerText)", "导出区出现（导入完成）", timeout=90)
            try:
                caps = wait_for(
                    ws,
                    "(() => {const b=[...document.querySelectorAll('button')].find(e=>e.textContent.trim()==='Excel');"
                    "return b && !b.disabled ? 'ready' : null;})()",
                    "report.formats 回包（Excel 按钮解除禁用）",
                    timeout=30,
                )
            except TimeoutError:
                # 禁用不等于"没数据"——必须把页面自己的报错行读出来，否则分不清
                # "内核没这个 RPC"、"RPC 竞态失败"和"数据没导入成功"
                diag_js = (
                    "(()=>{const t=document.body.innerText;"
                    "return {error:(t.match(/报告能力查询失败[^\n]*/)||'')[0]||null,"
                    " export:(t.match(/导出失败[^\n]*/)||'')[0]||null,"
                    " importing:/导入中…/.test(t), hasExportRow:/导出分析报告/.test(t)};})()"
                )
                print("诊断:", evaluate(ws, diag_js))
                raise
            print("Excel 按钮状态:", caps)

            word = evaluate(
                ws,
                "(() => {const b=[...document.querySelectorAll('button')].find(e=>e.textContent.trim()==='Word');"
                "return b ? {disabled: !!b.disabled, title: b.title || ''} : null;})()",
            )
            # 按钮状态只该由内核能力决定，所以这里按"跑 sidecar 的那个解释器里有没有
            # python-docx"来判——换台机器装了/没装，脚本都不用改，也不会给出假绿
            if args.packaged:
                # 装机版要看的是 resources 里那个冻结内核，它跟开发机装没装 python-docx 无关
                # （已单独用 dist/uhf-sidecar/uhf-sidecar.exe 实测 report.formats 含 word）
                expect_enabled = True
            else:
                import importlib.util

                expect_enabled = importlib.util.find_spec("docx") is not None
            print(f"Word 按钮: {word}  解释器有 docx={expect_enabled}")
            if not word or bool(word["disabled"]) == expect_enabled:
                failures.append(f"Word 按钮状态与内核能力不一致：{word}（应 enabled={expect_enabled}）")
            if not expect_enabled and "python-docx" not in word.get("title", ""):
                failures.append(f"禁用却没写明原因：{word}")

            # ── 导 Excel ──
            click_text(ws, "Excel")
            end = time.time() + 40
            while time.time() < end and not SAVE.exists():
                time.sleep(0.4)
            if not SAVE.exists():
                failures.append("点 Excel 后 PD_SAVE_PATH 没有落盘文件")
            else:
                time.sleep(1.0)
                dest = SHOTS / "report.xlsx"
                shutil.copy(SAVE, dest)
                SAVE.unlink()
                import openpyxl

                sheet = openpyxl.load_workbook(dest).active
                a1, b10 = sheet["A1"].value, sheet["B10"].value
                print(f"Excel 重新打开：A1={a1!r} 总事件数 B10={b10!r} 大小={dest.stat().st_size}B")
                if a1 != "PRPD 分析报告" or not isinstance(b10, int) or b10 <= 0:
                    failures.append(f"Excel 内容不对：A1={a1!r} B10={b10!r}")
                notice = evaluate(
                    ws,
                    "(()=>{const p=[...document.querySelectorAll('p')].find(e=>/已导出/.test(e.textContent));return p?p.textContent.trim():null;})()",
                )
                print("页面回执:", notice)
                if not notice or "已导出" not in notice:
                    failures.append(f"页面没有显示导出回执：{notice!r}")

            # ── 导 JSON ──
            click_text(ws, "JSON")
            end = time.time() + 40
            while time.time() < end and not SAVE.exists():
                time.sleep(0.4)
            if not SAVE.exists():
                failures.append("点 JSON 后没有落盘文件")
            else:
                time.sleep(1.0)
                dest = SHOTS / "report.json"
                shutil.copy(SAVE, dest)
                SAVE.unlink()
                parsed = json.loads(dest.read_text(encoding="utf-8"))
                total = parsed.get("statistics", {}).get("total_events")
                print(
                    f"JSON 重新解析：total_events={total} 分类={parsed.get('discharge_classification', {}).get('name')!r}"
                )
                if not isinstance(total, int) or total <= 0:
                    failures.append(f"JSON 统计不对：{total!r}")

            if expect_enabled:
                click_text(ws, "Word")
                end = time.time() + 40
                while time.time() < end and not SAVE.exists():
                    time.sleep(0.4)
                if not SAVE.exists():
                    failures.append("点 Word 后没有落盘文件")
                else:
                    time.sleep(1.0)
                    dest = SHOTS / "report.docx"
                    shutil.copy(SAVE, dest)
                    SAVE.unlink()
                    import docx

                    doc = docx.Document(dest)
                    shapes = len(doc.inline_shapes)
                    paragraphs = [par.text for par in doc.paragraphs]
                    has_chart_heading = "2. PRPD 图谱" in paragraphs
                    print(
                        f"Word 重新打开：段落 {len(doc.paragraphs)} 个，内嵌图 {shapes} 张，图谱小节={has_chart_heading}"
                    )
                    if shapes != 1 or not has_chart_heading:
                        failures.append(f"Word 里的图谱不对：图 {shapes} 张、小节={has_chart_heading}")
                    notice = evaluate(
                        ws,
                        "(()=>{const p=[...document.querySelectorAll('p')].find(e=>/已导出/.test(e.textContent));return p?p.textContent.trim():null;})()",
                    )
                    print("页面回执:", notice)
                    if not notice or "图谱" not in notice:
                        failures.append(f"回执没说明图谱是否嵌入：{notice!r}")

            err = evaluate(
                ws,
                "(()=>{const p=[...document.querySelectorAll('p')].find(e=>/导出失败|报告能力查询失败/.test(e.textContent));return p?p.textContent.trim():null;})()",
            )
            if err:
                failures.append(f"页面报错行：{err}")
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=15)
        except subprocess.TimeoutExpired:
            proc.kill()

    print("\nP2 真点验收:", "PASS" if not failures else "FAIL")
    for f in failures:
        print("  -", f)
    return 0 if not failures else 1


if __name__ == "__main__":
    sys.exit(main())
